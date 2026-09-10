from io import BytesIO

from docx import Document

from app.utils.document_parser import (
    extract_placeholder_keys,
    extract_variables,
    find_unresolved_placeholders,
    replace_placeholders,
)


def _split_across_runs(paragraph, text: str) -> None:
    """Simulate Word splitting a single token across several runs by
    adding one run per character."""
    for char in text:
        paragraph.add_run(char)


def test_extracts_variables_from_paragraphs_tables_and_headers() -> None:
    doc = Document()
    doc.add_paragraph("I, {{recipient_name}}, aged {{recipient_age}} years.")

    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].paragraphs[0].add_run("{{donor_name}}")
    table.rows[0].cells[1].paragraphs[0].add_run("{{blood_group}}")

    header = doc.sections[0].header
    header.paragraphs[0].add_run("Hospital: {{hospital_name}}")

    keys = extract_placeholder_keys(doc)

    assert keys == ["recipient_name", "recipient_age", "donor_name", "blood_group", "hospital_name"]


def test_ignores_blank_signature_lines_and_underscores() -> None:
    doc = Document()
    doc.add_paragraph("Signature: ____________________")
    doc.add_paragraph("Witness Name: {{witness_name}}")

    keys = extract_placeholder_keys(doc)

    assert keys == ["witness_name"]


def test_deduplicates_repeated_variable_across_document() -> None:
    doc = Document()
    doc.add_paragraph("{{donor_name}} on page 1.")
    doc.add_paragraph("{{donor_name}} again on page 5.")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].paragraphs[0].add_run("{{donor_name}} in a table too.")

    keys = extract_placeholder_keys(doc)

    assert keys == ["donor_name"]


def test_infers_field_types_from_key_name() -> None:
    doc = Document()
    doc.add_paragraph("{{recipient_age}} {{transplant_date}} {{recipient_address}} {{recipient_name}}")

    variables = {v.key: v for v in extract_variables(doc)}

    assert variables["recipient_age"].type == "number"
    assert variables["transplant_date"].type == "date"
    assert variables["recipient_address"].type == "textarea"
    assert variables["recipient_name"].type == "text"


def test_replaces_placeholder_split_across_multiple_runs() -> None:
    doc = Document()
    paragraph = doc.add_paragraph("I, ")
    _split_across_runs(paragraph, "{{recipient_name}}")
    paragraph.add_run(", hereby declare.")

    replace_placeholders(doc, {"recipient_name": "Jane Doe"})

    assert doc.paragraphs[0].text == "I, Jane Doe, hereby declare."
    assert find_unresolved_placeholders(doc) == []


def test_replaces_all_occurrences_of_a_repeated_variable() -> None:
    doc = Document()
    doc.add_paragraph("Donor: {{donor_name}}")
    doc.add_paragraph("Signed by {{donor_name}} below.")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].paragraphs[0].add_run("{{donor_name}}")

    replace_placeholders(doc, {"donor_name": "Yunus Mulani"})

    assert doc.paragraphs[0].text == "Donor: Yunus Mulani"
    assert doc.paragraphs[1].text == "Signed by Yunus Mulani below."
    assert table.rows[0].cells[0].paragraphs[0].text == "Yunus Mulani"


def test_replaces_inside_table_cells_without_breaking_table_layout() -> None:
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    headers = ["Name", "Age", "Relation"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.paragraphs[0].add_run(text)
    values = ["{{member_name}}", "{{member_age}}", "{{member_relation}}"]
    for cell, text in zip(table.rows[1].cells, values):
        cell.paragraphs[0].add_run(text)

    replace_placeholders(doc, {"member_name": "Hasina", "member_age": "45", "member_relation": "Wife"})

    assert len(table.rows) == 2
    assert len(table.columns) == 3
    assert table.rows[1].cells[0].paragraphs[0].text == "Hasina"
    assert table.rows[1].cells[1].paragraphs[0].text == "45"
    assert table.rows[1].cells[2].paragraphs[0].text == "Wife"


def test_round_trips_through_bytes_buffer() -> None:
    doc = Document()
    doc.add_paragraph("{{client_name}} confirms this affidavit.")
    replace_placeholders(doc, {"client_name": "Ramesh"})

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    reloaded = Document(buffer)
    assert reloaded.paragraphs[0].text == "Ramesh confirms this affidavit."
