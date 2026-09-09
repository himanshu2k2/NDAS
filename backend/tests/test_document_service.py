from io import BytesIO

import pytest
from docx import Document
from fastapi import HTTPException

from app.services.document_service import extract_template_variables, generate_affidavit_docx


def _make_docx(text: str) -> BytesIO:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def test_extract_template_variables_returns_specs_in_document_order() -> None:
    source = _make_docx("I, {{recipient_name}}, aged {{recipient_age}}, residing at {{recipient_address}}.")

    variables = extract_template_variables(source)

    assert [v.key for v in variables] == ["recipient_name", "recipient_age", "recipient_address"]
    assert variables[1].type == "number"
    assert variables[2].type == "textarea"


def test_extract_template_variables_rejects_non_docx_bytes() -> None:
    with pytest.raises(HTTPException) as exc_info:
        extract_template_variables(BytesIO(b"not a real docx file"))
    assert exc_info.value.status_code == 400


def test_generate_affidavit_docx_replaces_all_placeholders() -> None:
    source = _make_docx("Donor: {{donor_name}}, Recipient: {{recipient_name}}")

    result = generate_affidavit_docx(source, {"donor_name": "Yunus Mulani", "recipient_name": "Hasina Mulani"})

    generated = Document(result)
    assert generated.paragraphs[0].text == "Donor: Yunus Mulani, Recipient: Hasina Mulani"


def test_generate_affidavit_docx_raises_when_a_placeholder_is_left_unfilled() -> None:
    source = _make_docx("Donor: {{donor_name}}, Recipient: {{recipient_name}}")

    with pytest.raises(HTTPException) as exc_info:
        generate_affidavit_docx(source, {"donor_name": "Yunus Mulani"})

    assert exc_info.value.status_code == 422
    assert "recipient_name" in exc_info.value.detail
