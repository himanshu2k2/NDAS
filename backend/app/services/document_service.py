"""Template-driven DOCX generation.

Given a template file on disk and a dict of validated variable values,
produces a new DOCX in memory with every ``{{key}}`` placeholder replaced,
leaving the original template file untouched.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from fastapi import HTTPException

from app.utils.document_parser import extract_variables, find_unresolved_placeholders, replace_placeholders


def extract_template_variables(source) -> list:
    """Open a .docx (path or in-memory buffer) and return its extracted
    variable specs, or raise 400 if the file cannot be parsed as a Word
    document."""
    try:
        document = Document(source)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Uploaded file is not a valid Word (.docx) document.") from exc
    return extract_variables(document)


def generate_affidavit_docx(source, variable_data: dict) -> BytesIO:
    """Load the template (path or in-memory buffer), replace all
    placeholders, and return the completed document as an in-memory
    buffer. The source template is never modified."""
    try:
        document = Document(source)
    except Exception as exc:
        raise HTTPException(status_code=500, detail="Stored template could not be opened.") from exc

    string_values = {key: "" if value is None else str(value) for key, value in variable_data.items()}
    replace_placeholders(document, string_values)

    remaining = find_unresolved_placeholders(document)
    if remaining:
        raise HTTPException(
            status_code=422,
            detail=f"Missing values for placeholder(s): {', '.join(remaining)}",
        )

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
