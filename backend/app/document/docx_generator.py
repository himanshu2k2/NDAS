from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

from app.templates.kidney_affidavit_template import AffidavitContext, build_sections

_MARGIN = Inches(1.0)
_TITLE_SIZE = Pt(14)
_SUBTITLE_SIZE = Pt(10)
_BODY_SIZE = Pt(11)
_HEADING_SIZE = Pt(11)


def _set_margins(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.27)   # A4
        section.page_height = Inches(11.69)
        section.left_margin = _MARGIN
        section.right_margin = _MARGIN
        section.top_margin = _MARGIN
        section.bottom_margin = _MARGIN


def _add_page_break(doc: Document) -> None:
    para = doc.add_paragraph()
    run = para.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def _add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = _TITLE_SIZE
    run.font.color.rgb = RGBColor(0, 0, 0)


def _add_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = _SUBTITLE_SIZE


def _add_body_para(doc: Document, text: str) -> None:
    text = text.strip()
    if not text:
        doc.add_paragraph()
        return
    # Section headings (ALL CAPS lines like "PART I — ...")
    is_heading = text.isupper() or (
        text.startswith("PART ") and "—" in text
    ) or text in ("DEPONENT", "WITNESS", "NOTARY / OATH COMMISSIONER")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.size = _HEADING_SIZE if is_heading else _BODY_SIZE
    run.bold = is_heading
    if not is_heading:
        p.paragraph_format.first_line_indent = Inches(0.3)


def generate_docx(ctx: AffidavitContext) -> BytesIO:
    doc = Document()
    _set_margins(doc)

    # Remove default empty paragraph
    for para in doc.paragraphs:
        p = para._element
        p.getparent().remove(p)

    sections = build_sections(ctx)
    for i, section in enumerate(sections):
        if section["title"]:
            _add_title(doc, section["title"])
        if section.get("subtitle"):
            _add_subtitle(doc, section["subtitle"])
        doc.add_paragraph()  # spacer

        for para_text in section["body"]:
            _add_body_para(doc, para_text)

        if section["page_break_after"] and i < len(sections) - 1:
            _add_page_break(doc)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
